from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


WORKSPACE = Path(r"D:\越南MOM\MOM")
OUTPUT_DIR = WORKSPACE / "outputs"
OUTPUT_PATH = OUTPUT_DIR / "监视与测量设备管理_台账管理需求分析书_V1.0.docx"
SKILL_ROOT = Path(
    r"C:\Users\ptn_z\.codex\plugins\cache\openai-primary-runtime\documents\26.802.11031\skills\documents"
)
sys.path.insert(0, str(SKILL_ROOT / "scripts"))
from table_geometry import apply_table_geometry  # noqa: E402


BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "202124"
MUTED = "667085"
LIGHT_GRAY = "F2F4F7"
LIGHT_BLUE = "E8EEF5"
CALLOUT = "F4F6F9"
WHITE = "FFFFFF"
GREEN = "2F855A"
ORANGE = "B54708"
RED = "B42318"
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def set_font(run, size=None, color=None, bold=None, italic=None, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_paragraph_shading(paragraph, fill, border_color=None):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    if border_color:
        p_bdr = p_pr.find(qn("w:pBdr"))
        if p_bdr is None:
            p_bdr = OxmlElement("w:pBdr")
            p_pr.append(p_bdr)
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "18")
        left.set(qn("w:space"), "8")
        left.set(qn("w:color"), border_color)
        p_bdr.append(left)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph, value=True):
    paragraph.paragraph_format.keep_with_next = value


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = instruction
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr_text, fld_sep, text, fld_end])


def add_numbering_definition(doc, num_format, text, left=720, hanging=360):
    numbering = doc.part.numbering_part.element
    abstract_ids = [
        int(el.get(qn("w:abstractNumId")))
        for el in numbering.findall(qn("w:abstractNum"))
        if el.get(qn("w:abstractNumId")) is not None
    ]
    num_ids = [
        int(el.get(qn("w:numId")))
        for el in numbering.findall(qn("w:num"))
        if el.get(qn("w:numId")) is not None
    ]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), num_format)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), text)
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), str(left))
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), str(left))
    ind.set(qn("w:hanging"), str(hanging))
    p_pr.extend([tabs, ind])
    lvl.extend([start, fmt, lvl_text, lvl_jc, p_pr])
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_num(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_el])


def add_bullet(doc, text, num_id):
    p = doc.add_paragraph()
    apply_num(p, num_id)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    set_font(p.add_run(text), size=11, color=INK)
    return p


def add_step(doc, text, num_id):
    return add_bullet(doc, text, num_id)


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.1
    if bold_lead and text.startswith(bold_lead):
        set_font(p.add_run(bold_lead), size=11, color=INK, bold=True)
        set_font(p.add_run(text[len(bold_lead):]), size=11, color=INK)
    else:
        set_font(p.add_run(text), size=11, color=INK)
    return p


def add_label_value(doc, label, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.1
    set_font(p.add_run(label + "："), size=10.5, color=INK, bold=True)
    set_font(p.add_run(value), size=10.5, color=INK)
    return p


def add_callout(doc, label, text, color=BLUE):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Twips(220)
    p.paragraph_format.right_indent = Twips(120)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.15
    set_paragraph_shading(p, CALLOUT, color)
    set_font(p.add_run(label + "  "), size=10.5, color=color, bold=True)
    set_font(p.add_run(text), size=10.5, color=INK)
    return p


def add_table(doc, headers, rows, widths, header_fill=LIGHT_GRAY, font_size=9.2, alignments=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, header in enumerate(headers):
        cell = hdr.cells[idx]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, header_fill)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        set_font(p.add_run(str(header)), size=9.2, color=INK, bold=True)

    for row_data in rows:
        row = table.add_row()
        for idx, value in enumerate(row_data):
            cell = row.cells[idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            if alignments and idx < len(alignments):
                p.alignment = alignments[idx]
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_font(p.add_run(str(value)), size=font_size, color=INK)
    apply_table_geometry(
        table,
        widths,
        table_width_dxa=CONTENT_WIDTH_DXA,
        indent_dxa=TABLE_INDENT_DXA,
        cell_margins_dxa={"top": 80, "bottom": 80, "start": 120, "end": 120},
    )
    after = doc.add_paragraph()
    after.paragraph_format.space_before = Pt(0)
    after.paragraph_format.space_after = Pt(2)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    set_keep_with_next(p)
    return p


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    configs = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for style_name, (size, color, before, after) in configs.items():
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True


def configure_page(doc):
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.right_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)


def configure_header_footer(section):
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    set_font(hp.add_run("MOM · 监视与测量设备管理"), size=8.5, color=MUTED, bold=True)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(0)
    set_font(fp.add_run("台账管理需求分析书  |  第 "), size=8.5, color=MUTED)
    add_field(fp, " PAGE ")
    set_font(fp.add_run(" 页"), size=8.5, color=MUTED)


def page_break(doc):
    doc.add_page_break()


def build_document():
    doc = Document()
    configure_styles(doc)
    configure_page(doc)
    for section in doc.sections:
        configure_header_footer(section)

    bullet_num = add_numbering_definition(doc, "bullet", "•")
    decimal_num = add_numbering_definition(doc, "decimal", "%1.")

    # Cover: standard_business_brief with memo_masthead title override (23 pt, black).
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(18)
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_after = Pt(4)
    set_font(kicker.add_run("需求分析书"), size=10.5, color=BLUE, bold=True)
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(5)
    set_font(title.add_run("监视与测量设备管理"), size=23, color="000000", bold=True)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(20)
    set_font(subtitle.add_run("台账管理子模块"), size=15, color=MUTED, bold=True)

    add_label_value(doc, "文件编号", "MOM-EQM-LDG-RA-001")
    add_label_value(doc, "版本", "V1.0（评审稿）")
    add_label_value(doc, "编制日期", "2026-08-03")
    add_label_value(doc, "编制单位", "MOM 项目组")
    add_label_value(doc, "适用对象", "产品、质量/计量、设备、生产使用部门、IT 开发与测试")
    add_callout(
        doc,
        "文档定位",
        "本文件以当前台账管理原型为业务基线，定义正式上线所需的功能、数据、规则、权限和验收标准。原型中的演示数据与提示信息不代表后台接口、数据库及业务闭环已经完成。",
    )
    page_break(doc)

    add_heading(doc, "1. 文档说明", 1)
    add_heading(doc, "1.1 编制目的", 2)
    add_body(doc, "统一监视与测量设备台账管理的业务口径，为产品评审、开发设计、接口联调、测试验收及上线培训提供可追溯依据。")
    add_heading(doc, "1.2 文档基线", 2)
    add_table(
        doc,
        ["基线项", "说明"],
        [
            ["系统范围", "MOM 系统 > 设备管理 > 监视与测量设备管理 > 台账管理"],
            ["页面基线", "LedgerManagement.vue；包含列表、查询、新增、导入、导出、列选择及批量操作入口"],
            ["关联页面", "校验项目管理、校验模板管理、校验任务管理、校验提醒管理"],
            ["数据现状", "当前页面采用前端示例数据，部分按钮仅展示消息或尚未绑定业务处理"],
            ["需求状态", "V1.0 评审稿；待第 15 章事项确认后形成开发基线"],
        ],
        [2100, 7260],
        font_size=9.5,
    )
    add_heading(doc, "1.3 版本记录", 2)
    add_table(
        doc,
        ["版本", "日期", "变更内容", "状态"],
        [["V1.0", "2026-08-03", "首次形成台账管理需求分析，补充列选择与校验业务闭环", "评审稿"]],
        [1100, 1700, 4960, 1600],
        alignments=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER],
    )

    add_heading(doc, "2. 业务背景与建设目标", 1)
    add_heading(doc, "2.1 业务背景", 2)
    add_body(doc, "监视与测量设备包括量具、检具及其他用于过程监视、检验和计量的设备。设备数量增长后，若依赖分散台账或人工提醒，容易出现编号重复、责任不清、临期遗漏、过期仍使用、证书难追溯等问题。")
    add_body(doc, "当前页面已呈现较完整的台账字段，但字段数量多，固定全部展示会降低操作效率。因此需要允许操作者按岗位选择展示列，同时保证关键业务字段、校验任务和设备生命周期信息保持一致。")
    add_heading(doc, "2.2 建设目标", 2)
    for item in [
        "建立唯一、完整、可追溯的监视与测量设备主台账。",
        "打通登记、使用、临期提醒、校验任务、校验结果回写、封存与报废的生命周期。",
        "降低多字段列表的使用负担，支持岗位化列展示并保留用户偏好。",
        "通过权限、数据范围和操作日志控制关键数据变更风险。",
        "支持批量导入、导出和批量维护，提升存量数据治理效率。",
    ]:
        add_bullet(doc, item, bullet_num)
    add_heading(doc, "2.3 成功判定", 2)
    add_callout(doc, "上线判定", "P0 需求全部通过验收；设备管理编号无重复；临期/过期计算与校验任务状态一致；关键变更可通过日志追溯到人员、时间和前后值。", GREEN)

    add_heading(doc, "3. 范围与边界", 1)
    add_heading(doc, "3.1 本期范围", 2)
    for item in [
        "台账列表、分页、组合查询、列选择及个人展示偏好。",
        "台账新增、详情、编辑、附件、批量导入、导出和受控批量修改。",
        "从台账选择设备生成校验任务，并接收校验完成后的结果回写。",
        "设备正常、封存、报废等使用状态管理，以及临期、校验中、过期等校验状态展示。",
        "操作权限、数据范围、异常提示、操作日志及基础性能要求。",
    ]:
        add_bullet(doc, item, bullet_num)
    add_heading(doc, "3.2 非本期范围", 2)
    for item in [
        "校验项目、校验模板内部的完整编辑与审批规则（本文件仅定义台账引用约束）。",
        "自动采集计量仪器读数、与第三方计量机构的在线送检接口。",
        "固定资产折旧、采购、费用结算及财务资产管理。",
        "二维码标签打印、移动端盘点与离线作业（可作为后续扩展）。",
    ]:
        add_bullet(doc, item, bullet_num)

    add_heading(doc, "4. 用户角色与权限", 1)
    add_body(doc, "权限分为“操作权限”和“数据范围”两层。操作权限控制能做什么，数据范围控制能看到和维护哪些部门/车间的数据。")
    add_table(
        doc,
        ["角色", "主要职责", "建议权限", "数据范围"],
        [
            ["计量管理员", "维护台账、组织校验、处理证书与状态", "查询、新增、编辑、导入导出、生成任务、封存/报废", "全组织或授权工厂"],
            ["部门设备责任人", "维护本部门设备位置和责任信息，发起校验", "查询、有限编辑、生成任务、查看证书", "所属部门/车间"],
            ["校验执行人员", "执行内部或外部校验任务并回填结果", "查询、查看详情、关联任务", "被分配任务涉及设备"],
            ["质量/审核人员", "监督有效期、抽查校验证据和变更记录", "查询、导出、查看日志与附件", "授权组织"],
            ["系统管理员", "配置权限、字典、接口和异常处理", "系统配置及必要的维护权限", "全局；业务修改需审计"],
            ["普通查看者", "查询设备有效状态和责任信息", "只读查询", "所属组织"],
        ],
        [1400, 2500, 3260, 2200],
        font_size=8.8,
    )
    add_callout(doc, "业务保护", "已产生校验任务或历史记录的设备不得通过普通“删除”清除；应转为封存、报废或停用，并保留完整追溯链。", RED)

    add_heading(doc, "5. 核心业务流程", 1)
    add_heading(doc, "5.1 设备生命周期主流程", 2)
    steps = [
        "登记设备：录入基础信息、计量方式、校验模板、周期、责任部门及附件。",
        "校验数据校验：系统检查管理编号唯一性、必填项、日期关系及组织归属。",
        "投入使用：设备使用状态为“正常”，系统根据有效日期计算校验状态与预警天数。",
        "临期提醒：达到提醒规则时进入“临期”，提示责任人并可自动或人工生成校验任务。",
        "执行校验：校验任务进入“校验中”，台账禁止重复生成未关闭的同类任务。",
        "结果回写：任务完成后回写校验日期、有效日期、证书编号、校验单位和结论。",
        "后续处置：合格则恢复有效；不合格则按规则进入维修、停用、封存或报废处理。",
    ]
    for step in steps:
        add_step(doc, step, decimal_num)
    add_heading(doc, "5.2 状态模型", 2)
    add_body(doc, "建议将页面当前单一“状态”拆分为两个维度，避免把设备使用状态与校验有效状态混在一起。")
    add_table(
        doc,
        ["状态维度", "状态值", "产生方式", "业务含义"],
        [
            ["设备使用状态", "正常 / 封存 / 报废", "人工受控操作", "决定设备是否允许被选用或生成新任务"],
            ["校验状态", "有效 / 临期 / 校验中 / 过期", "系统计算 + 任务状态", "反映设备计量有效性，不替代使用状态"],
            ["校验结论", "合格 / 不合格 / 待判定", "校验任务完成回写", "决定设备能否恢复正常使用"],
        ],
        [1700, 2200, 2400, 3060],
        font_size=9.2,
    )
    add_heading(doc, "5.3 关键异常流程", 2)
    for item in [
        "管理编号重复：阻止保存或导入，并指明冲突行和已有设备。",
        "设备过期且无任务：高亮标识；禁止用于受控生产，允许授权人员立即生成任务。",
        "存在未关闭任务：再次生成时提示已有任务编号，不重复创建。",
        "校验不合格：不得自动恢复“有效”；必须记录处置意见和责任人。",
        "委外校验：允许先创建送检任务，回厂后补录证书及校验结果。",
    ]:
        add_bullet(doc, item, bullet_num)

    add_heading(doc, "6. 功能需求", 1)
    add_heading(doc, "6.1 功能清单", 2)
    requirements = [
        ["FR-LDG-001", "台账列表与分页", "P0", "按权限范围展示设备；选择、序号和操作列固定保留。"],
        ["FR-LDG-002", "快速/高级查询", "P0", "支持关键字及用途、编号、日期、周期、方式、状态、组织、预警天数组合查询。"],
        ["FR-LDG-003", "列选择", "P0", "27 个业务字段可选；支持全选、全部显示、恢复默认，并保存个人偏好。"],
        ["FR-LDG-004", "新增台账", "P0", "校验必填、唯一性、日期和组织信息；保存后形成可追溯记录。"],
        ["FR-LDG-005", "详情与编辑", "P0", "详情展示完整字段、附件、校验摘要和变更记录；编辑受权限及状态控制。"],
        ["FR-LDG-006", "删除/停用控制", "P0", "无引用草稿可删除；有历史数据只能封存或报废。"],
        ["FR-LDG-007", "批量导入", "P0", "模板下载、预校验、错误报告、确认导入及导入结果汇总。"],
        ["FR-LDG-008", "导出", "P1", "按当前查询范围导出，可选择当前展示列或全部业务字段。"],
        ["FR-LDG-009", "批量修改", "P1", "仅开放责任人、组织、位置等白名单字段，并二次确认。"],
        ["FR-LDG-010", "生成校验任务", "P0", "从所选设备生成任务；校验适用性并防止重复任务。"],
        ["FR-LDG-011", "封存/报废", "P0", "记录原因、日期、操作者及附件；变更后禁止生成普通校验任务。"],
        ["FR-LDG-012", "证书与附件", "P1", "上传、预览、下载、版本保留，并关联设备或校验任务。"],
        ["FR-LDG-013", "提醒联动", "P0", "按计量方式和提醒规则计算临期状态，支持自动生成任务。"],
        ["FR-LDG-014", "操作日志", "P0", "记录新增、编辑、导入、任务生成、状态变更和关键导出。"],
        ["FR-LDG-015", "权限与数据范围", "P0", "操作按钮和数据查询均按角色、组织及管理权限控制。"],
    ]
    add_table(doc, ["需求编号", "功能", "优先级", "摘要"], requirements, [1650, 2100, 1050, 4560], font_size=8.7,
              alignments=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT])

    add_heading(doc, "6.2 列表、查询与分页", 2)
    for item in [
        "默认按最近更新时间倒序；分页默认 10 条，可切换 20/50 条。",
        "快速查询匹配设备名称和管理编号；按 Enter 或“查询”执行。",
        "高级查询条件可组合；日期区间包含起止日期；重置恢复全部条件。",
        "查询条件、分页和列选择不改变后台数据，仅影响当前用户视图。",
        "列表无数据时显示明确空状态；接口失败时保留查询条件并提供重试。",
    ]:
        add_bullet(doc, item, bullet_num)

    add_heading(doc, "6.3 列选择", 2)
    add_body(doc, "列选择用于解决台账字段过多的问题。选择框、序号和操作列不进入选择范围；其余 27 个业务字段由操作者自由配置。")
    add_table(
        doc,
        ["项目", "需求"],
        [
            ["默认展示 13 列", "量具类型、管理编号、名称、规格型号、量程、计量周期/月、校验日期、有效日期、状态、使用部门、使用车间、责任人、预警天数"],
            ["可选字段", "当前台账定义的 27 个业务字段；字段名称与表头保持一致"],
            ["快捷操作", "全选/取消全选、全部显示、恢复默认；显示已选数量"],
            ["即时生效", "勾选或取消后列表立即增减对应列，不要求重新查询"],
            ["偏好保存", "P0 可先按浏览器本地保存；P1 建议绑定用户账号，以便跨设备同步"],
            ["导出联动", "导出时默认使用当前展示列，同时提供“全部字段”选项"],
        ],
        [2100, 7260],
        font_size=9.4,
    )

    add_heading(doc, "6.4 新增、编辑与详情", 2)
    for item in [
        "新增时至少校验管理编号、名称、生产厂家、计量周期、校验日期和计量方式。",
        "管理编号保存后普通用户不可修改；确需变更时走专用更码权限并保留旧编号映射。",
        "编辑页面显示字段级校验信息；保存成功后更新版本、修改人和修改时间。",
        "详情页展示基础信息、组织位置、当前校验信息、证书附件、关联任务和变更历史。",
        "封存、报废或校验中的设备，字段可编辑范围由业务状态和权限共同控制。",
    ]:
        add_bullet(doc, item, bullet_num)

    add_heading(doc, "6.5 导入、导出与批量修改", 2)
    add_body(doc, "批量能力用于数据治理，不应绕过单条维护规则。所有批量操作必须先校验、再预览、后确认，并记录结果。")
    add_table(
        doc,
        ["功能", "处理步骤", "失败处理"],
        [
            ["批量导入", "下载模板 → 上传 → 格式/字典/唯一性校验 → 预览 → 确认导入", "输出包含行号、字段、错误原因的失败明细；错误行不落库"],
            ["导出", "按查询结果与用户选择字段生成 Excel；记录操作者、条件、条数和时间", "异步任务失败可重试；不得导出权限范围外数据"],
            ["批量修改", "选择设备 → 选择允许字段 → 输入新值 → 影响预览 → 二次确认", "部分失败时返回逐条结果；不允许静默跳过"],
        ],
        [1600, 4100, 3660],
        font_size=9.0,
    )
    add_callout(doc, "批量字段白名单", "建议首期仅开放使用部门、使用车间、班组、具体位置、责任人和备注。管理编号、校验日期、有效日期、证书编号及状态不得通过普通批量修改。", ORANGE)

    add_heading(doc, "6.6 生成校验任务", 2)
    for item in [
        "至少选择一台设备；系统逐台校验使用状态、校验方式、模板和现有未关闭任务。",
        "内部校验设备必须绑定已确认的适用模板；委外校验设备至少具备校验方式、责任人和计划完成日期。",
        "同一设备不得同时存在两个未关闭的同类校验任务。",
        "生成成功后返回任务编号清单，并可跳转至校验任务管理。",
        "校验任务完成后，由统一服务事务性回写台账校验信息，避免任务与台账不一致。",
    ]:
        add_bullet(doc, item, bullet_num)

    add_heading(doc, "6.7 封存与报废", 2)
    add_body(doc, "封存用于暂时停止使用且仍保留恢复可能的设备；报废为不可逆的业务终态。两者均必须填写原因和日期，建议要求附件或审批依据。")
    add_table(
        doc,
        ["操作", "前置条件", "结果", "恢复规则"],
        [
            ["封存", "无执行中的校验任务，或先关闭/取消任务", "使用状态=封存；不可生成普通校验任务", "授权人员填写启封原因后恢复正常"],
            ["报废", "确认无在用关联；完成必要审批", "使用状态=报废；保留历史和附件", "原则上不可恢复；纠错需管理员审批并留痕"],
            ["删除", "仅无业务引用的草稿/误建记录", "物理或逻辑删除，记录操作日志", "按系统策略处理，普通用户不可恢复"],
        ],
        [1300, 3000, 2960, 2100],
        font_size=8.9,
    )

    add_heading(doc, "7. 数据需求", 1)
    add_heading(doc, "7.1 台账业务字段字典", 2)
    fields = [
        ["量具类型", "字典/50", "否", "长度、质量、温度、专用检具等；建议字典化"],
        ["用途", "枚举/20", "否", "量具、检具；可扩展"],
        ["管理类别", "枚举/10", "否", "A/B/C；业务含义需由计量管理制度确认"],
        ["校验模板", "引用", "条件必填", "内校必选已确认且适用的模板；外校按规则配置"],
        ["管理编号", "字符串/50", "是", "组织范围或全局唯一；建议支持编码规则生成"],
        ["出厂编号", "字符串/100", "否", "制造商序列号；允许无编号但需说明"],
        ["入厂日期", "日期", "否", "不得晚于当前日期，历史补录需权限"],
        ["名称", "字符串/100", "是", "设备标准名称"],
        ["规格型号", "字符串/100", "否", "支持字母、数字及常用规格符号"],
        ["量程", "字符串/100", "否", "首期文本；后续可拆最小值、最大值和单位"],
        ["分度值", "字符串/50", "否", "首期文本；应与量程单位一致"],
        ["材质", "字符串/50", "否", "设备或检具主要材质"],
        ["生产厂家", "字符串/150", "是", "制造商名称"],
        ["计量周期/月", "正整数", "是", "默认字典 3/6/12 月；允许受控扩展"],
        ["校验日期", "日期", "是", "最近一次有效校验日期"],
        ["有效日期", "日期", "系统计算", "按校验日期与周期计算；人工调整需专用权限"],
        ["校验单位", "字符串/150", "否", "内校部门或外部计量机构"],
        ["证书编号", "字符串/100", "否", "校验完成后回写；同证书复用规则待确认"],
        ["计量方式", "枚举/20", "是", "内校、外校"],
        ["状态", "组合展示", "系统/人工", "建议拆为设备使用状态与校验状态"],
        ["使用部门", "组织引用", "建议必填", "受当前用户管理权限限制"],
        ["使用车间", "组织引用", "否", "必须属于所选使用部门"],
        ["班组", "组织引用", "否", "必须属于所选车间或部门"],
        ["具体位置", "字符串/150", "否", "检验台、货架、区域等"],
        ["责任人", "用户引用", "建议必填", "应为有效用户且在授权组织范围内"],
        ["预警天数", "计算字段", "系统计算", "有效日期减当前日期；过期显示负值或“已过期”"],
        ["备注", "字符串/500", "否", "补充说明，不承载关键状态"],
        ["附件", "文件集合", "否", "图片/PDF/Office；记录文件名、版本、上传人和时间"],
    ]
    add_table(doc, ["字段", "类型/长度", "必填", "规则/来源"], fields, [1800, 1550, 1150, 4860], font_size=8.5,
              alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT])

    add_heading(doc, "7.2 系统审计字段", 2)
    add_table(
        doc,
        ["字段", "说明"],
        [
            ["主键 ID", "系统唯一标识，不直接供业务人员修改"],
            ["数据版本号", "用于乐观锁，防止多人编辑互相覆盖"],
            ["创建人/创建时间", "首次入账信息"],
            ["修改人/修改时间", "最近一次修改信息"],
            ["删除标识", "逻辑删除控制；有业务引用时不得删除"],
            ["组织/工厂标识", "数据隔离和权限过滤依据"],
        ],
        [2300, 7060],
        font_size=9.5,
    )

    add_heading(doc, "7.3 核心实体关系", 2)
    for item in [
        "设备台账 1:N 校验任务：每次校验形成独立任务，任务保留当时的设备关键快照。",
        "设备台账 N:1 校验模板：设备引用当前适用模板，历史任务不随模板变更。",
        "设备台账 1:N 附件/证书：附件可属于设备档案或具体校验任务。",
        "设备台账 1:N 操作日志：记录关键字段前值、后值、操作者、时间、来源和原因。",
        "用户 1:1 或 1:N 列偏好：建议按用户、页面和组织保存展示字段配置。",
    ]:
        add_bullet(doc, item, bullet_num)

    add_heading(doc, "8. 业务规则", 1)
    rules = [
        ["BR-001", "唯一性", "管理编号在约定范围内唯一；导入和接口保存均需校验。"],
        ["BR-002", "日期", "校验日期不得晚于有效日期；有效日期计算口径必须全系统一致。"],
        ["BR-003", "预警", "预警天数=有效日期-业务当前日期；提醒规则按计量方式或设备类别匹配。"],
        ["BR-004", "任务幂等", "同一设备存在未关闭校验任务时不得重复创建。"],
        ["BR-005", "模板适用", "内校任务只能引用已确认且适用该设备的模板。"],
        ["BR-006", "结果回写", "只有已完成且结论明确的任务可更新台账最近校验信息。"],
        ["BR-007", "不合格控制", "不合格设备不得自动恢复有效使用，必须记录处置。"],
        ["BR-008", "历史保留", "有校验历史的设备不得物理删除；封存/报废不清除历史。"],
        ["BR-009", "并发控制", "编辑提交时校验数据版本，冲突时提示重新加载。"],
        ["BR-010", "批量安全", "批量操作逐条校验、结果可追踪，不因部分失败造成整体状态不明。"],
        ["BR-011", "权限一致", "前端按钮隐藏不能替代后端鉴权；查询与导出使用同一数据范围。"],
        ["BR-012", "列偏好", "无偏好时使用默认 13 列；偏好失效字段自动剔除并保留其余有效配置。"],
    ]
    add_table(doc, ["规则编号", "主题", "规则描述"], rules, [1500, 1600, 6260], font_size=9.1)

    add_heading(doc, "9. 页面与交互要求", 1)
    add_heading(doc, "9.1 页面布局", 2)
    for item in [
        "顶部左侧为标题、快速查询、查询/重置和高级查询；右侧为新增及批量操作。",
        "高级查询按业务分组展示，展开后不挤压表格操作列；页面高度变化时表格区域自适应。",
        "列选择采用紧凑弹层，多列排列并限制最大高度，字段过多时弹层内部滚动。",
        "表格保留横向滚动，操作列固定在右侧；长文本使用省略及悬浮提示。",
        "危险操作使用醒目样式，并在执行前展示对象数量、影响和不可逆提示。",
    ]:
        add_bullet(doc, item, bullet_num)
    add_heading(doc, "9.2 反馈与异常", 2)
    add_table(
        doc,
        ["场景", "反馈要求"],
        [
            ["查询中", "显示加载状态，避免重复请求；失败后保留原列表并提示重试"],
            ["保存成功", "明确提示成功并刷新当前记录；保持当前查询与分页位置"],
            ["表单校验失败", "定位到具体字段，提示可执行的修正信息"],
            ["批量操作", "展示处理中、成功数、失败数和可下载失败明细"],
            ["无权限", "提示无权限或数据范围限制，不暴露越权对象详情"],
            ["并发冲突", "提示数据已被他人修改，可重新加载后再编辑"],
        ],
        [2200, 7160],
        font_size=9.4,
    )

    add_heading(doc, "10. 接口与集成要求", 1)
    add_body(doc, "以下定义逻辑服务能力，不限定最终 URL 或技术框架。接口需统一返回业务码、消息、请求标识和字段级错误。")
    add_table(
        doc,
        ["服务能力", "主要输入", "主要输出/约束"],
        [
            ["台账分页查询", "查询条件、分页、排序、组织范围", "设备摘要、总数；后端强制数据权限过滤"],
            ["台账新增/编辑", "设备字段、版本号、附件引用", "设备 ID、版本号；唯一性和并发校验"],
            ["台账详情", "设备 ID", "完整字段、校验摘要、附件、关联任务和操作历史"],
            ["导入校验/确认", "模板文件、导入批次", "预校验结果、错误报告、导入结果"],
            ["导出任务", "查询条件、字段清单", "异步任务编号、下载状态；复用查询权限"],
            ["生成校验任务", "设备 ID 列表、计划日期、方式", "逐台成功/失败及任务编号；具备幂等控制"],
            ["封存/报废", "设备 ID、版本号、原因、日期、附件", "更新后状态及审计编号"],
            ["列偏好", "用户、页面、字段清单", "保存/读取个人展示配置"],
        ],
        [2100, 3100, 4160],
        font_size=8.9,
    )
    add_heading(doc, "10.1 关联基础数据", 2)
    for item in [
        "组织、部门、车间、班组及用户来自系统基础信息，禁止自由文本形成重复组织。",
        "校验模板和校验项目来自已确认版本；台账只保存引用标识和必要快照。",
        "提醒规则由校验提醒管理维护，台账展示计算结果，不重复维护规则。",
        "编码规则可与系统编码规则配置集成，是否自动生成由第 15 章确认。",
    ]:
        add_bullet(doc, item, bullet_num)

    add_heading(doc, "11. 非功能需求", 1)
    nfrs = [
        ["性能", "常规分页查询在 95% 场景下 2 秒内返回；单页 50 条；导入/导出采用异步任务。"],
        ["容量", "首期至少支持 10 万条设备台账和长期校验历史，附件独立存储。"],
        ["安全", "后端鉴权、组织数据隔离、文件类型校验、下载鉴权和敏感操作二次确认。"],
        ["审计", "关键操作日志不可由普通业务用户修改；保留期按公司制度配置。"],
        ["可用性", "查询、编辑和批量操作提供明确状态；浏览器刷新不丢失已保存列偏好。"],
        ["一致性", "任务完成与台账回写采用事务或可靠消息，失败可重试并可人工补偿。"],
        ["兼容性", "支持项目约定的主流桌面浏览器分辨率；表格横向滚动不遮挡操作列。"],
        ["国际化", "字段文案遵循系统中/英/越语言机制；编码、证书号和用户输入按 Unicode 存储。"],
        ["备份恢复", "台账、任务、日志和附件索引纳入备份；恢复演练满足系统总体 RPO/RTO。"],
    ]
    add_table(doc, ["类别", "要求"], nfrs, [1700, 7660], font_size=9.2)

    add_heading(doc, "12. 原型现状与实施差距", 1)
    gaps = [
        ["列表与查询", "已有前端示例数据和本地过滤", "接入分页查询接口、权限过滤、排序和异常处理"],
        ["列选择", "27 个字段可选，默认 13 列，本地浏览器保存", "根据评审决定是否增加账号级服务端同步"],
        ["新增", "已有表单和部分必填校验，提交后仅提示成功", "接入唯一性校验、保存接口、附件上传和刷新"],
        ["详情/编辑/删除", "页面有操作入口，未形成完整处理", "实现详情、受控编辑、删除保护及审计"],
        ["批量导入", "已有上传弹窗和模板下载提示", "实现模板文件、预校验、错误报告和确认入库"],
        ["导出", "当前仅提示任务已提交", "实现异步导出、下载中心及权限一致性"],
        ["生成校验任务", "已有按钮入口", "接入模板校验、幂等创建、结果清单和任务跳转"],
        ["批量修改/报废封存", "已有按钮入口", "实现白名单、二次确认、审批/原因及日志"],
        ["数据存储", "示例数据位于前端", "建设数据库、服务接口、附件存储和数据迁移"],
    ]
    add_table(doc, ["能力", "当前原型", "正式上线需完成"], gaps, [1900, 3000, 4460], font_size=8.8)
    add_callout(doc, "实施原则", "开发评估不得把现有按钮或成功提示当作后端功能完成证据。每项 P0 能力必须具备接口、数据持久化、权限、异常处理和验收用例。", RED)

    add_heading(doc, "13. 验收标准与测试场景", 1)
    acceptance = [
        ["AC-001", "默认列", "首次进入显示约定 13 个业务字段，选择/序号/操作列固定。", "P0"],
        ["AC-002", "列偏好", "增减字段即时生效；刷新后配置仍保留；恢复默认正确。", "P0"],
        ["AC-003", "组合查询", "多条件组合结果准确，重置后恢复全部数据。", "P0"],
        ["AC-004", "新增唯一性", "重复管理编号无法保存，并指出冲突对象。", "P0"],
        ["AC-005", "日期规则", "非法日期关系被阻止；有效日期按确认口径计算。", "P0"],
        ["AC-006", "权限范围", "不同角色只能看到和操作授权组织数据；导出范围一致。", "P0"],
        ["AC-007", "导入校验", "错误文件生成逐行错误报告；错误记录不落库。", "P0"],
        ["AC-008", "导出", "导出行数、过滤条件和字段与用户选择一致。", "P1"],
        ["AC-009", "任务生成", "符合条件设备生成任务；重复未关闭任务被拦截。", "P0"],
        ["AC-010", "任务回写", "任务完成后台账日期、证书、结论和状态一致更新。", "P0"],
        ["AC-011", "不合格处置", "不合格设备不能自动恢复有效，必须记录处置。", "P0"],
        ["AC-012", "封存", "封存后不可生成普通任务，历史记录仍可查看。", "P0"],
        ["AC-013", "报废", "报废操作包含原因和日志，普通用户不能恢复。", "P0"],
        ["AC-014", "批量修改", "仅白名单字段可选；返回逐条成功/失败结果。", "P1"],
        ["AC-015", "并发编辑", "版本冲突时不覆盖他人修改，并提示重新加载。", "P0"],
        ["AC-016", "审计", "关键操作可查到操作者、时间、对象及前后值。", "P0"],
        ["AC-017", "性能", "常规分页查询达到第 11 章响应目标。", "P0"],
        ["AC-018", "多语言", "切换语言后页面主要字段和操作文案按系统机制展示。", "P1"],
    ]
    add_table(doc, ["编号", "场景", "通过标准", "级别"], acceptance, [1350, 1700, 5260, 1050], font_size=8.6,
              alignments=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER])

    add_heading(doc, "14. 数据迁移与上线建议", 1)
    rollout = [
        "盘点现有 Excel/纸质台账，统一管理编号、组织、状态和日期口径。",
        "建立字典及组织映射，先执行全量预校验，输出重复编号和缺失必填项。",
        "选择一个部门试点导入，核对查询、提醒、任务生成和证书回写。",
        "修复数据问题后分批迁移；每批记录来源文件、批次、成功数和失败数。",
        "上线前冻结旧台账的新增入口，完成差异补录与业务签字确认。",
        "上线后保留短期回查窗口，按设备数量、临期数、过期数和任务闭环率核对。",
    ]
    for item in rollout:
        add_step(doc, item, decimal_num)
    add_heading(doc, "14.1 建议实施优先级", 2)
    add_table(
        doc,
        ["阶段", "范围", "退出条件"],
        [
            ["阶段一：台账可用", "查询、列选择、新增/编辑、详情、权限、数据库", "台账数据可持久化且权限正确"],
            ["阶段二：业务闭环", "提醒、生成校验任务、结果回写、封存/报废、日志", "生命周期 P0 用例全部通过"],
            ["阶段三：效率提升", "导入导出、批量修改、账号级列偏好、性能优化", "批量与体验类用例通过"],
        ],
        [1900, 4160, 3300],
        font_size=9.0,
    )

    add_heading(doc, "15. 待确认事项", 1)
    decisions = [
        ["D-01", "管理编号是人工输入、按编码规则自动生成，还是两者兼容？", "产品/计量/IT", "开发前"],
        ["D-02", "管理编号唯一范围是集团、工厂还是组织？", "计量/质量", "开发前"],
        ["D-03", "有效日期采用自然月、固定天数还是减一天口径？", "计量/质量", "开发前"],
        ["D-04", "管理类别 A/B/C 的定义及对周期、审批的影响？", "计量/质量", "开发前"],
        ["D-05", "封存、启封和报废是否需要审批；审批人和流程是什么？", "设备/质量", "开发前"],
        ["D-06", "校验不合格后的维修、复检和让步使用流程如何处理？", "质量/生产", "联调前"],
        ["D-07", "内校与外校分别要求哪些模板、证书和环境信息？", "计量", "开发前"],
        ["D-08", "附件允许的格式、单文件大小、数量和保留年限？", "IT/质量", "开发前"],
        ["D-09", "提醒渠道仅站内，还是包含邮件、企业消息或看板？", "产品/IT", "联调前"],
        ["D-10", "列偏好首期仅本地保存，还是直接实现账号跨设备同步？", "产品/IT", "开发前"],
        ["D-11", "普通用户是否允许删除无引用记录，还是统一只允许管理员删除？", "质量/设备", "开发前"],
        ["D-12", "导出是否默认当前展示列；是否需要导出审批或水印？", "质量/信息安全", "联调前"],
    ]
    add_table(doc, ["编号", "待确认问题", "建议责任方", "确认节点"], decisions, [1200, 5000, 1900, 1260], font_size=8.8,
              alignments=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER])
    add_callout(doc, "评审结论要求", "第 15 章事项确认后，应更新本文件版本并冻结 P0 需求、字段口径、状态模型和接口边界；未确认项不得由开发自行假设。", BLUE)

    add_heading(doc, "附录 A：默认列与全部可选列", 1)
    add_heading(doc, "A.1 固定列", 2)
    add_body(doc, "选择框、序号、操作。固定列不参与列选择，确保批量操作和单条维护入口始终可用。")
    add_heading(doc, "A.2 默认业务列", 2)
    add_body(doc, "量具类型、管理编号、名称、规格型号、量程、计量周期/月、校验日期、有效日期、状态、使用部门、使用车间、责任人、预警天数。")
    add_heading(doc, "A.3 全部可选业务列", 2)
    optional_rows = [
        ["基础分类", "量具类型、用途、管理类别、校验模板"],
        ["标识与入厂", "管理编号、出厂编号、入厂日期"],
        ["规格属性", "名称、规格型号、量程、分度值、材质、生产厂家"],
        ["校验信息", "计量周期/月、校验日期、有效日期、校验单位、证书编号、计量方式、状态"],
        ["使用信息", "使用部门、使用车间、班组、具体位置、责任人"],
        ["辅助信息", "预警天数、备注"],
    ]
    add_table(doc, ["分组", "字段"], optional_rows, [1900, 7460], font_size=9.5)

    add_heading(doc, "附录 B：需求追溯说明", 1)
    add_body(doc, "功能需求编号 FR-LDG-001 至 FR-LDG-015 与验收场景 AC-001 至 AC-018 共同构成首期需求追溯基线。开发任务、接口任务和测试用例应引用对应编号，避免仅按页面按钮拆分导致业务规则遗漏。")
    add_body(doc, "若评审后调整状态模型、有效日期计算或删除策略，必须同步更新字段字典、业务规则、接口约束和验收用例。")

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    output = build_document()
    print(output)
