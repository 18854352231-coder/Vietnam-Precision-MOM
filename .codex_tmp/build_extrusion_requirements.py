from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"D:\越南MOM\MOM")
OUT = ROOT / "docs" / "挤压生产报表需求分析书.docx"

BLUE = "2E74B5"
DEEP_BLUE = "1F4D78"
TEXT = "202124"
MUTED = "667085"
LIGHT = "F2F4F7"
PALE_BLUE = "EAF2F8"
PALE_GREEN = "EAF7EF"
PALE_AMBER = "FFF4E5"
PALE_RED = "FDECEC"
BORDER = "D0D5DD"
WHITE = "FFFFFF"

CONTENT_DXA = 9360


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    assert sum(widths) == CONTENT_DXA, (sum(widths), widths)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    cell_mar = tbl_pr.find(qn("w:tblCellMar"))
    if cell_mar is None:
        cell_mar = OxmlElement("w:tblCellMar")
        tbl_pr.append(cell_mar)
    for side, amount in (("top", 80), ("bottom", 80), ("start", 120), ("end", 120)):
        node = cell_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            cell_mar.append(node)
        node.set(qn("w:w"), str(amount))
        node.set(qn("w:type"), "dxa")

    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), BORDER)

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_text(cell, text: str, *, bold=False, color=TEXT, size=8.5, align=None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    if align is not None:
        p.alignment = align
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int], *, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    header = table.rows[0]
    for idx, text in enumerate(headers):
        set_cell_text(header.cells[idx], text, bold=True, color=DEEP_BLUE, size=font_size, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(header.cells[idx], LIGHT)
    set_repeat_table_header(header)
    set_cant_split(header)

    for row_values in rows:
        row = table.add_row()
        set_cant_split(row)
        for idx, text in enumerate(row_values):
            set_cell_text(row.cells[idx], text, size=font_size)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_status_table(doc: Document, rows: list[list[str]]):
    table = add_table(doc, ["级别", "判定", "本报表含义"], rows, [1500, 2100, 5760], font_size=9)
    fills = {"A": PALE_GREEN, "B": PALE_BLUE, "C": PALE_AMBER, "D": PALE_RED}
    for row in table.rows[1:]:
        set_cell_shading(row.cells[0], fills.get(row.cells[0].text, WHITE))
        row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    return table


def add_label_paragraph(doc: Document, label: str, text: str, *, fill=PALE_BLUE):
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    set_cell_text(table.cell(0, 0), label, bold=True, color=DEEP_BLUE, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_shading(table.cell(0, 0), fill)
    set_cell_text(table.cell(0, 1), text, size=9)
    set_repeat_table_header(table.rows[0])
    set_cant_split(table.rows[0])
    set_table_geometry(table, [1500, 7860])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_bullet(doc: Document, text: str, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.add_run(text)
    return p


def add_number(doc: Document, text: str):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)
    return p


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)
    paragraph.add_run(" 页")


def add_heading(doc: Document, text: str, level: int):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def build_document() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name in ("List Bullet", "List Bullet 2", "List Number"):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167

    heading_specs = {
        "Title": (23, DEEP_BLUE, 0, 8),
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DEEP_BLUE, 8, 4),
    }
    for style_name, (size, color, before, after) in heading_specs.items():
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    p = header.paragraphs[0]
    p.text = "MOM 生产报表  |  挤压生产报表需求分析"
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in p.runs:
        run.font.name = "Calibri"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    add_page_number(fp)
    for run in fp.runs:
        run.font.name = "Calibri"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MUTED)

    # Memo masthead
    eyebrow = doc.add_paragraph()
    eyebrow.paragraph_format.space_after = Pt(5)
    r = eyebrow.add_run("内部需求分析  ·  MOM / 挤压与锯切联合报表")
    r.bold = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string(BLUE)

    title = doc.add_paragraph(style="Title")
    title.add_run("挤压生产报表需求分析书")
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    r = subtitle.add_run("字段口径、数据来源、计算规则与落库要求")
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor.from_string(MUTED)

    meta = doc.add_table(rows=2, cols=4)
    meta.style = "Table Grid"
    meta_rows = [
        ["文档版本", "V1.0", "编制日期", "2026-08-12"],
        ["适用模块", "挤压与锯切联合报表", "依据", "当前前端代码"],
    ]
    for i, values in enumerate(meta_rows):
        for j, value in enumerate(values):
            set_cell_text(meta.cell(i, j), value, bold=(j % 2 == 0), color=DEEP_BLUE if j % 2 == 0 else TEXT, size=9)
            if j % 2 == 0:
                set_cell_shading(meta.cell(i, j), LIGHT)
    set_repeat_table_header(meta.rows[0])
    for row in meta.rows:
        set_cant_split(row)
    set_table_geometry(meta, [1300, 2600, 1300, 4160])

    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(10)
    rule.paragraph_format.space_after = Pt(10)
    p_pr = rule._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), BLUE)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)

    add_label_paragraph(
        doc,
        "结论先行",
        "当前联合报表只展示已完工数据：挤压记录须达到“已完成”，锯切记录须达到“已完成”。锯切数据已由锯切工作台写入浏览器 localStorage；挤压数据仍为报表内置演示数据，工作台中的挤压数量、批次和剔料状态也未形成稳定的跨页面持久化记录。正式上线前必须补齐挤压排程同步、生产事件落库和后端查询接口。",
        fill=PALE_AMBER,
    )

    add_heading(doc, "1. 文档目的与范围", 1)
    doc.add_paragraph(
        "本分析书定义“挤压与锯切联合报表”的业务目标、数据粒度、查询条件、展示字段、获取方式、计算公式、关联规则、数据质量要求和验收标准。所有实现状态均以 2026-08-12 的当前代码为依据；原型中的静态数组和浏览器本地存储不等同于生产数据库已接通。"
    )
    add_heading(doc, "1.1 建设目标", 2)
    add_bullet(doc, "按生产排程查看挤压投入、挤压执行和锯切执行结果，形成一行贯通的生产追溯视图。")
    add_bullet(doc, "只保留能够从前序业务记录直接取得，或能由可靠明细计算得到的字段。")
    add_bullet(doc, "统一屏幕查询、表格展示和 Excel 导出的字段口径，禁止同名指标使用不同算法。")
    add_bullet(doc, "明确原型可用数据与正式上线缺口，为后端表结构、接口和事件采集提供依据。")

    add_heading(doc, "1.2 本期范围", 2)
    add_table(
        doc,
        ["范围项", "纳入内容", "不纳入内容"],
        [
            ["报表对象", "挤压排程、挤压生产记录、锯切生产记录", "设备 OEE、能耗、工艺参数趋势"],
            ["展示层", "查询条件、联合明细、分页、Excel 导出、汇总指标", "自定义拖拽列、图表驾驶舱"],
            ["数据层", "字段来源、关联键、计算规则、持久化要求", "本需求书内直接设计数据库物理索引"],
            ["已删除字段", "不再展示完成进度、锯切实际完成、完工框数", "不恢复已确认删除字段"],
        ],
        [1700, 4100, 3560],
        font_size=9,
    )

    add_heading(doc, "2. 当前实现盘点", 1)
    add_status_table(
        doc,
        [
            ["A", "可直接获取且已持久化", "字段已有稳定记录，可直接查询；当前主要为锯切汇总记录。"],
            ["B", "可直接获取但未持久化", "页面或工作台已有值，但刷新、切换浏览器或跨设备后不能保证。"],
            ["C", "可计算但输入未稳定落库", "公式明确，必须先持久化参与计算的明细或事件。"],
            ["D", "当前不可可靠获取", "现有对象缺少关联键、历史记录或可靠语义，不能作为正式报表值。"],
        ],
    )

    add_heading(doc, "2.1 代码现状", 2)
    add_table(
        doc,
        ["模块", "当前数据形态", "结论"],
        [
            ["挤压报表", "ExtrusionReport.vue 内 sourceRows 为 9 条静态演示数据，其中仅满足完工条件的数据进入报表。", "不能视为已接通挤压生产数据。"],
            ["挤压排程", "IssueScheduling.vue 已有排程字段，但“下发”只更新页面状态；未同步到挤压工作台共享存储。", "排程基础字段可定义来源，但需要接口或共享数据层。"],
            ["挤压工作台", "已上料队列、已挤压状态、挤压批次和剔料列表主要保存在组件内存。", "字段有业务动作来源，但刷新后不可靠。"],
            ["锯切工作台", "开工、装框、取样、报废和完工会汇总为 SawingProductionRecord。", "公式已形成，并写入本地记录。"],
            ["锯切存储", "mom_sawing_production_records 保存于当前浏览器 localStorage。", "可支撑原型联调，不能替代服务端数据库。"],
        ],
        [1600, 4760, 3000],
        font_size=9,
    )

    add_heading(doc, "2.2 目标数据流", 2)
    add_number(doc, "排程下发时，以排程编号生成或同步一条挤压排程主记录。")
    add_number(doc, "上料、挤压、剔料、退料等动作写入带排程编号和挤压批次号的生产事件。")
    add_number(doc, "锯切开工、装框、取样、报废、完工动作写入锯切生产记录或事件明细。")
    add_number(doc, "报表服务先限定已完工记录，再按查询条件读取排程和生产记录，在服务端完成聚合与关联。")
    add_number(doc, "页面与 Excel 导出复用同一查询结果和同一字段口径。")

    add_heading(doc, "3. 报表粒度与关联规则", 1)
    add_label_paragraph(doc, "数据粒度", "默认一行代表一个生产排程；如同一排程允许多个挤压批次或多个锯台并行，则生产实现必须改为“排程编号 + 挤压批次号 + 锯切作业号”粒度，不能覆盖记录。")
    add_label_paragraph(doc, "准入规则", "挤压数据仅当挤压状态=已完成时进入报表；锯切数据仅当锯切状态=已完成时进入报表。未开始、生产中或未完工记录不得参与页面、汇总和导出。", fill=PALE_GREEN)
    add_heading(doc, "3.1 联合规则", 2)
    add_table(
        doc,
        ["优先级", "关联条件", "使用要求", "异常处理"],
        [
            ["1", "挤压.scheduleNo = 锯切.scheduleNo", "排程编号必须全局唯一且在工序间透传。", "重复排程编号进入数据质量告警。"],
            ["2", "orderNo + extrusionBatchNo", "仅当两个字段均非空时允许兜底关联。", "任一为空时不得关联，防止空批次误匹配。"],
            ["3", "锯切独立记录", "锯切有记录但未匹配挤压时仍展示。", "挤压专属字段显示“—”，不得显示为 0。"],
        ],
        [950, 2600, 3450, 2360],
        font_size=9,
    )
    add_label_paragraph(doc, "实现提醒", "当前前端兜底逻辑未限制 extrusionBatchNo 非空；正式实现必须增加非空校验，并优先使用不可变的排程主键或作业主键。", fill=PALE_RED)

    add_heading(doc, "4. 查询条件需求", 1)
    filter_rows = [
        ["计划日期", "日期区间", "排程.planDate；锯切独立记录取锯切.planDate", "闭区间查询：startDate ≤ planDate ≤ endDate", "B/A"],
        ["排程编号", "手动输入", "排程.scheduleNo / 锯切.scheduleNo", "去首尾空格，不区分大小写，包含匹配", "B/A"],
        ["订单编号", "手动输入", "排程.orderNo / 锯切.orderNo", "去首尾空格，不区分大小写，包含匹配", "B/A"],
        ["产品名称", "手动输入", "排程.productName / 锯切.productName", "去首尾空格，不区分大小写，包含匹配", "B/A"],
        ["挤压批次号", "手动输入", "挤压生产批次 / 锯切.extrusionBatchNo", "包含匹配；正式接口建议精确查询可选", "B/A"],
        ["合金牌号", "手动输入", "排程.alloy / 锯切.alloy", "当前已改为手动输入，包含匹配", "B/A"],
    ]
    add_table(doc, ["字段", "控件", "数据来源", "查询规则", "成熟度"], filter_rows, [1100, 1100, 2700, 3300, 1160], font_size=8.3)
    add_bullet(doc, "点击“查询”后才应用草稿条件；点击“重置”清空所有条件并回到第一页。")
    add_bullet(doc, "筛选、分页、汇总和导出必须基于同一已完工结果集；导出不受当前分页限制。")

    add_heading(doc, "5. 明细字段与数据获取方式", 1)
    doc.add_paragraph("成熟度按第 2 章定义。目标表/对象名称为逻辑建议，最终可映射到现有数据库命名规范。")

    add_heading(doc, "5.1 排程信息", 2)
    rows = [
        ["序号", "页面计算", "无", "(当前页码−1)×每页条数+行内索引+1", "A"],
        ["计划日期", "直接获取", "ExtrusionSchedule.planDate；锯切独立行取 SawingProductionRecord.planDate", "YYYY-MM-DD；优先取挤压排程日期", "B/A"],
        ["排程编号", "直接获取", "ExtrusionSchedule.scheduleNo / SawingProductionRecord.scheduleNo", "联合主键；必填、唯一、不可变", "B/A"],
        ["订单编号", "直接获取", "ExtrusionSchedule.orderNo；锯切兜底取同名字段", "原值展示；可参与非空兜底关联", "B/A"],
        ["产品名称", "直接获取", "ExtrusionSchedule.productName；锯切兜底取同名字段", "原值展示", "B/A"],
    ]
    add_table(doc, ["字段", "方式", "来源", "口径/公式", "成熟度"], rows, [950, 1100, 3250, 3000, 1060], font_size=8.3)

    add_heading(doc, "5.2 物料与模具", 2)
    rows = [
        ["合金牌号", "直接获取", "ExtrusionSchedule.alloy；锯切独立行取 SawingProductionRecord.alloy", "以排程主数据为准；两工序不一致时告警", "B/A"],
        ["铝棒料号", "直接获取", "ExtrusionSchedule.rodMaterialNo", "锯切独立行无挤压记录时显示“—”", "B"],
        ["棒径(mm)", "直接获取", "ExtrusionSchedule.rodDiameter", "数值，单位 mm；缺失显示“—”", "B"],
        ["棒长(mm)", "直接获取", "ExtrusionSchedule.rodLength", "数值，单位 mm；缺失显示“—”", "B"],
        ["模具组号", "直接/计算", "ExtrusionSchedule.moldGroupNo；缺失时可由 moldNo 推导", "moldGroupNo 优先；否则去除模具号最后一个“-”分段，例如 M10-0649-200 → M10-0649", "B/C"],
        ["模具号", "直接获取", "ExtrusionSchedule.moldNo；锯切独立行取 SawingProductionRecord.moldNo", "多模并行时需拆分作业粒度或建立模具明细", "B/A"],
    ]
    add_table(doc, ["字段", "方式", "来源", "口径/公式", "成熟度"], rows, [950, 1100, 3250, 3000, 1060], font_size=8.3)

    add_heading(doc, "5.3 挤压执行", 2)
    rows = [
        ["挤压批次号", "直接获取", "上模/开工时生成的 ExtrusionProduction.extrusionBatchNo", "必须与排程编号、模具号同时落库；当前仅在工作台对象和队列内存中", "B"],
        ["计划支数", "直接获取", "ExtrusionSchedule.inputQtyRods 或下发后的 issueQty", "取下发时快照，禁止随主数据修改而追溯变动", "B"],
        ["已上料", "事件聚合", "ExtrusionRodEvent，eventType=FEED", "COUNT(有效上料事件)；退料应冲销。若仅统计当前队列，则字段应改名“当前在制棒数”", "C"],
        ["已挤压", "事件聚合", "ExtrusionRodEvent，eventType=EXTRUDE 或 rod.status=已挤压", "COUNT(完成挤压且未作废的棒号)；按排程/批次去重", "C"],
        ["待挤压", "计算", "已上料、已挤压", "MAX(已上料−已挤压, 0)", "C"],
        ["剔料", "事件聚合", "ExtrusionRodEvent，eventType=REJECT", "COUNT(有效剔料事件−撤销/重新投入事件)；必须带 scheduleNo", "D→C"],
        ["挤压状态", "计算", "计划支数、已上料、已挤压", "已挤压≥计划支数且计划>0：已完成；否则已上料>0或已挤压>0：生产中；否则：未开始", "C"],
    ]
    add_table(doc, ["字段", "方式", "来源", "口径/公式", "成熟度"], rows, [950, 1100, 3250, 3000, 1060], font_size=8.1)
    add_label_paragraph(doc, "状态风险", "若剔料导致最终合格挤压数小于计划支数，现有状态公式可能长期停留在“生产中”。正式业务需明确完成条件是“合格数达计划”还是“排程人工完工”，建议增加排程完工事件作为最终状态依据。", fill=PALE_AMBER)

    add_heading(doc, "5.4 锯切执行", 2)
    rows = [
        ["锯台", "直接获取", "SawingProductionRecord.machineNo", "锯切开工或切换锯台时取 currentMachine；正式记录应保存设备主键", "A"],
        ["定尺长度(mm)", "直接获取", "SawingSchedule.fixedLength → SawingProductionRecord.fixedLength", "数值，单位 mm；按排程/作业快照保存", "A"],
        ["装框数量", "计算", "frameListData.branchQty", "SUM(各料框支数)；当前对应 framedQty", "A*"],
        ["料框数", "计算", "frameListData", "COUNT(料框记录)；包括未完工框还是仅有效框需按当前定义为全部框", "A*"],
        ["取样数", "计算", "sampleListData", "COUNT(有效取样记录)", "A*"],
        ["报废数", "计算", "scrapHistoryData.scrapQty", "SUM(有效报废记录数量)", "A*"],
        ["开始时间", "直接获取", "SawingProductionRecord.startTime", "确认开工时写入；后续编辑不得覆盖首次开工时间", "A"],
        ["结束时间", "直接获取", "SawingProductionRecord.endTime", "确认完工时写入；未完工显示“—”", "A"],
        ["锯切状态", "直接获取", "SawingProductionRecord.status", "开工及过程变更为“生产中”；确认完工为“已完成”", "A"],
    ]
    add_table(doc, ["字段", "方式", "来源", "口径/公式", "成熟度"], rows, [950, 1100, 3250, 3000, 1060], font_size=8.1)
    doc.add_paragraph("注：A* 表示计算结果已写入 localStorage 汇总记录，但其明细列表当前仍是前端原型对象；生产化后应按排程隔离并由服务端持久化。")

    add_heading(doc, "6. 汇总指标与导出", 1)
    add_heading(doc, "6.1 汇总指标", 2)
    rows = [
        ["联合排程数", "COUNT(筛选后的已完工联合行)", "包含已完工锯切独立行；应按最终报表粒度去重"],
        ["计划支数", "SUM(有挤压记录行.planQty)", "锯切独立行不计入，避免重复或语义不一致"],
        ["已上料支数", "SUM(有挤压记录行.loadedQty)", "需先完成挤压事件持久化"],
        ["已挤压支数", "SUM(有挤压记录行.extrudedQty)", "需按排程/批次去重"],
        ["当前剔料数", "SUM(有挤压记录行.rejectedQty)", "建议改名“有效剔料数”，口径为未撤销的剔料事件"],
    ]
    add_table(doc, ["指标", "计算公式", "规则"], rows, [1900, 3000, 4460], font_size=9)

    add_heading(doc, "6.2 Excel 导出", 2)
    add_bullet(doc, "工作表“指标汇总”：导出上述 5 个指标、数值和单位。")
    add_bullet(doc, "工作表“挤压锯切联合明细”：只导出已完工结果，并包含第 5 章全部当前展示字段，字段顺序与页面一致。")
    add_bullet(doc, "已确认删除的“完成进度、锯切实际完成、完工框数”不得出现在页面、汇总或导出中。")
    add_bullet(doc, "导出文件名需包含导出日期；建议增加查询条件和导出时间元数据，便于追溯。")

    add_heading(doc, "7. 建议数据模型与接口", 1)
    add_heading(doc, "7.1 逻辑数据对象", 2)
    rows = [
        ["ExtrusionSchedule", "scheduleId、scheduleNo、orderNo、planDate、productName、alloy、rodMaterialNo、rodDiameter、rodLength、moldGroupNo、moldNo、planQty", "排程下发时形成不可变快照"],
        ["ExtrusionProduction", "scheduleId、extrusionBatchNo、moldNo、startTime、endTime、status、updatedAt", "一排程多批次时一对多"],
        ["ExtrusionRodEvent", "eventId、scheduleId、extrusionBatchNo、rodNo、eventType、qty、eventTime、operator、isReversed", "上料、挤压、剔料、退料均写事件"],
        ["SawingProduction", "sawingJobId、scheduleId、extrusionBatchNo、machineId、fixedLength、startTime、endTime、status、updatedAt", "避免只以排程编号覆盖多锯切作业"],
        ["SawingFrame", "sawingJobId、frameNo、branchQty、status", "计算装框数量和料框数"],
        ["SawingSample", "sawingJobId、sampleCode、sampleType、samplingTime、isValid", "计算取样数"],
        ["SawingScrap", "sawingJobId、scrapReason、scrapQty、scrapTime、isReversed", "计算报废数"],
    ]
    add_table(doc, ["对象", "核心字段", "说明"], rows, [1800, 4800, 2760], font_size=8.4)

    add_heading(doc, "7.2 查询接口建议", 2)
    add_table(
        doc,
        ["项目", "要求"],
        [
            ["接口", "GET /api/reports/extrusion-sawing（名称可按现有规范调整）"],
            ["入参", "startDate、endDate、scheduleNo、orderNo、productName、extrusionBatchNo、alloy、page、pageSize；服务端固定限制已完工"],
            ["返回", "items、total、summary、queryTime；items 直接返回页面所需展示字段和 hasExtrusion/hasSawing 标记"],
            ["排序", "默认 planDate 倒序、scheduleNo 倒序；相同排程多作业时按 startTime 倒序"],
            ["一致性", "items 与 summary 必须在同一查询快照或同一事务视图内计算"],
            ["权限", "按组织、车间、产线或设备数据权限过滤；导出权限单独控制"],
        ],
        [1800, 7560],
        font_size=9,
    )

    add_heading(doc, "8. 数据质量与显示规则", 1)
    rules = [
        ["空值", "未知或未关联显示“—”；只有业务明确为零时显示 0。"],
        ["数量", "所有数量为非负整数；负值仅允许作为事件冲销，不直接显示。"],
        ["时间", "统一 YYYY-MM-DD HH:mm:ss；数据库保存时区信息，页面按工厂时区展示。"],
        ["单位", "棒径、棒长、定尺长度统一 mm；数量统一支/条，字段名或表头明确单位。"],
        ["去重", "棒号、料框号、样品号在各自业务粒度内唯一；聚合前过滤作废和撤销记录。"],
        ["冲突", "排程与锯切的订单、产品、合金、模具不一致时，不静默覆盖；记录异常并以排程快照为主。"],
        ["更新", "生产动作成功提交后刷新报表数据；后端以 updatedAt 或版本号解决并发覆盖。"],
        ["追溯", "保留操作人、操作时间、来源设备、撤销原因；汇总值必须可下钻到明细事件。"],
    ]
    add_table(doc, ["主题", "规则"], rules, [1600, 7760], font_size=9)

    add_heading(doc, "9. 非功能需求", 1)
    add_table(
        doc,
        ["类别", "要求"],
        [
            ["性能", "默认 31 天范围、10 条/页，普通查询建议 3 秒内返回；导出大数据量采用异步任务。"],
            ["可靠性", "生产数据以服务端为唯一事实来源，不依赖 localStorage；写入失败必须提示且可重试。"],
            ["安全", "接口鉴权、数据权限、导出审计；不得在前端硬编码敏感生产数据。"],
            ["可维护性", "字段口径集中定义；页面、接口和 Excel 共享 DTO/字段映射，避免重复计算。"],
            ["兼容性", "支持当前项目目标浏览器；超宽表格保持横向滚动和固定关键列。"],
        ],
        [1600, 7760],
        font_size=9,
    )

    add_heading(doc, "10. 验收标准", 1)
    criteria = [
        ["AC-01", "排程下发后，报表可通过排程编号查询到基础字段，刷新页面后数据仍存在。"],
        ["AC-02", "完成上料、挤压、剔料和撤销动作后，已上料、已挤压、待挤压、剔料按本书公式更新。"],
        ["AC-03", "锯切开工、装框、取样、报废过程不进入报表；确认完工后，相应字段可刷新并跨浏览器/设备查询。"],
        ["AC-04", "排程编号优先关联；只有订单编号和挤压批次号同时非空时才允许兜底关联。"],
        ["AC-05", "锯切独立记录可显示，缺少挤压来源的字段展示“—”而不是 0。"],
        ["AC-06", "计划日期、排程编号、订单编号、产品名称、挤压批次号、合金牌号均能正确组合查询；查询区不显示执行状态条件。"],
        ["AC-07", "页面、汇总和 Excel 的同名字段数值一致；导出覆盖全部筛选结果而非当前页。"],
        ["AC-08", "页面和导出均不出现完成进度、锯切实际完成、完工框数。"],
        ["AC-09", "未开始、生产中或未完工记录不出现在页面、汇总和导出；无数据、空值、异常关联和重复事件按第 8 章处理。"],
        ["AC-10", "报表数据可追溯到排程、生产作业及具体事件，能定位操作时间与操作人。"],
    ]
    add_table(doc, ["编号", "验收条件"], criteria, [1300, 8060], font_size=9)

    add_heading(doc, "11. 实施优先级", 1)
    add_table(
        doc,
        ["优先级", "工作项", "交付结果"],
        [
            ["P0", "挤压排程下发同步；挤压批次、上料、挤压、剔料事件服务端落库；联合查询接口", "报表核心字段真实可用"],
            ["P0", "修正兜底关联的非空限制；确定一排程多批次/多锯台的数据粒度", "避免错配和记录覆盖"],
            ["P1", "锯切 localStorage 迁移至后端；装框、取样、报废明细按作业隔离", "跨终端共享且可追溯"],
            ["P1", "统一汇总与导出服务；补充权限、审计和数据质量告警", "正式业务可控上线"],
            ["P2", "异常下钻、异步大数据导出、性能优化", "提升分析与运维能力"],
        ],
        [1200, 5000, 3160],
        font_size=9,
    )

    add_heading(doc, "12. 代码依据", 1)
    sources = [
        ["联合报表字段、筛选、合并、汇总与导出", "src/views/reports/extrusion/ExtrusionReport.vue", "1–180、203–445"],
        ["挤压排程字段与下发动作", "src/views/extrusion/IssueScheduling.vue", "419–443、1086–1090"],
        ["挤压上料队列、批次、挤压与剔料状态", "src/views/extrusion/ExtrusionWorkbench.vue", "571–705、1286–1314、1365–1385"],
        ["锯切排程、开工、完工和汇总写入", "src/views/extrusion/SawingWorkbench.vue", "524–615、795–819、1196–1242"],
        ["锯切生产记录结构与本地存储", "src/utils/sawingProductionFlow.ts", "1–49"],
    ]
    add_table(doc, ["依据内容", "文件", "关键行"], sources, [3300, 4200, 1860], font_size=8.8)

    add_label_paragraph(doc, "文档边界", "本书对“字段从哪里来、怎样计算、当前是否可靠”作出定义；不代表后端接口、数据库表和跨终端数据同步已经实现。上线验收应以真实业务数据回放和接口联调结果为准。", fill=PALE_AMBER)

    core = doc.core_properties
    core.title = "挤压生产报表需求分析书"
    core.subject = "挤压与锯切联合报表字段口径、数据来源与计算规则"
    core.author = "Codex"
    core.keywords = "MOM, 挤压报表, 锯切报表, 需求分析, 字段来源, 计算公式"
    core.comments = "基于 2026-08-12 当前前端代码整理"

    doc.save(OUT)
    print(str(OUT))


if __name__ == "__main__":
    build_document()
