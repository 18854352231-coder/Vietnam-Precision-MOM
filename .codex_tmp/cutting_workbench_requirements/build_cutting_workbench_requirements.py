from __future__ import annotations

import hashlib
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


REFERENCE = Path(r"C:\Users\ptn_z\Desktop\挤压生产报表需求分析书.docx")
OUTPUT = Path(r"D:\越南MOM\MOM\docs\裁切工作台需求分析书.docx")
EXPECTED_SHA256 = "e6450c96a87602b130f010963bf4e170f1e5640a54c86eca72f287a716de5a1e"

HEADER_FILL = "F2F4F7"
HEADER_TEXT = "1F4D78"
BODY_TEXT = "202124"
BORDER = "D0D5DD"


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def clear_body(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_font(run, size=11, bold=None, color=BODY_TEXT) -> None:
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_text(cell, text: str, *, header=False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    if header:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(str(text))
    set_font(run, size=8, bold=header, color=HEADER_TEXT if header else BODY_TEXT)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if header:
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), HEADER_FILL)
        tc_pr.append(shd)


def set_cell_width(cell, width: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    total = sum(widths)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    cell_mar = tbl_pr.find(qn("w:tblCellMar"))
    if cell_mar is None:
        cell_mar = OxmlElement("w:tblCellMar")
        tbl_pr.append(cell_mar)
    for side, amount in (("top", 80), ("bottom", 80), ("start", 120), ("end", 120)):
        node = cell_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            cell_mar.append(node)
        node.set(qn("w:w"), str(amount))
        node.set(qn("w:type"), "dxa")

    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), BORDER)

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])


def mark_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, header=True)
    mark_header(table.rows[0])
    for values in rows:
        row = table.add_row()
        for index, value in enumerate(values):
            set_cell_text(row.cells[index], value)
    set_table_geometry(table, widths)
    doc.add_paragraph()


def add_heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)


def add_number(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)


def build() -> None:
    if sha256(REFERENCE) != EXPECTED_SHA256:
        raise RuntimeError("Reference DOCX changed; re-distillation is required")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(REFERENCE, OUTPUT)
    doc = Document(OUTPUT)
    clear_body(doc)

    title = doc.add_paragraph(style="Title")
    title.add_run("裁切工作台需求分析书")

    add_heading(doc, "文档目的与范围", 1)
    doc.add_paragraph(
        "按照已下发裁切排程和来源料框数据，完成班组上班、料框收料、裁切报工、装托装箱、标识卡打印及下工序流转；所有生产数量均须能够由前序记录直接取得或由可靠明细计算得到。"
    )

    add_heading(doc, "1.1 建设目标", 2)
    add_bullet(doc, "贯通裁切排程、来源料框、收料、裁切报工、长支扫码、装托装箱和标识打印，形成可追溯的工作台操作链路。")
    add_bullet(doc, "只保留可从排程、料框、扫码及报工记录直接获取，或可按明确公式计算得到的字段。")
    add_bullet(doc, "关键动作必须记录排程编号、来源料框、机台、班组、操作时间和操作人，支持后续裁切报表和包装工序使用。")

    add_heading(doc, "1.2 目标数据流", 2)
    add_number(doc, "排程下发后，以排程编号同步裁切排程主数据及来源料框明细。")
    add_number(doc, "操作人员选择裁切机台和班组完成上班，形成机台、班组和上班时间记录。")
    add_number(doc, "物料清单按排程载入来源料框；选择待收料且已有裁切排程的料框执行上料，记录收料状态和上料时间。")
    add_number(doc, "裁切完成后按料框报工，登记本次裁切数量、不良数量和不良原因，计算累计良品、不良品并记录完工时间。")
    add_number(doc, "选择排程和已收料料框，扫描长支镭雕码；需要装箱时先生成小箱码，再按箱容量绑定长支。")
    add_number(doc, "打印前生成栈板号和二维码预览；确认完成栈板后生成装托记录，并同步到后续料框列表。")
    add_number(doc, "正式上线时所有业务记录由服务端持久化；浏览器 localStorage 和页面内存仅用于当前原型联调。")

    add_heading(doc, "2. 查询条件需求", 1)
    query_rows = [
        ["物料清单-框号", "手动输入", "去首尾空格后按框号包含匹配"],
        ["物料清单-挤压批次", "手动输入", "按挤压批次号包含匹配"],
        ["物料清单-炉次号", "手动输入", "按炉次号包含匹配"],
        ["物料清单-模具号", "手动输入", "按模具号包含匹配"],
        ["物料清单-状态", "下拉选择", "待收料、已收料、已完工精确匹配"],
        ["装托记录-排程编号", "手动输入", "按排程编号包含匹配"],
        ["装托记录-栈板编号", "手动输入", "优先匹配 palletNo，缺失时匹配 frameNo"],
        ["装托记录-挤压批次", "手动输入", "按挤压批次号包含匹配"],
        ["装托记录-客户名称", "手动输入", "按客户名称包含匹配"],
    ]
    add_table(doc, ["字段", "控件", "查询规则"], query_rows, [1581, 1113, 5215])
    add_bullet(doc, "点击“重置”清空当前标签页查询条件；查询结果应与当前工作台最新状态一致。")
    add_bullet(doc, "正式实现应由服务端完成筛选和分页，禁止只筛选浏览器内的演示数组。")

    add_heading(doc, "3. 明细字段与数据获取方式", 1)

    add_heading(doc, "3.1 上下班与设备", 2)
    add_table(doc, ["字段", "方式"], [
        ["上班状态", "操作状态"],
        ["裁切机台", "手动选择"],
        ["上班班组", "手动选择"],
        ["上班时间", "系统生成"],
        ["下班时间", "系统生成（目标）"],
        ["操作人", "登录用户获取（目标）"],
    ], [1538, 2160])

    add_heading(doc, "3.2 裁切排程", 2)
    add_table(doc, ["字段", "方式"], [
        ["排程编号", "直接获取"],
        ["排程类型", "直接获取"],
        ["客户代码", "直接获取"],
        ["客户名称", "直接获取"],
        ["炉次号", "直接获取/多值合并"],
        ["挤压批次号", "直接获取/多值合并"],
        ["模具号", "直接获取/多值合并"],
        ["合金牌号", "直接获取"],
        ["产品名称", "直接获取"],
        ["组件物料号", "直接获取"],
        ["客户物料号", "直接获取"],
        ["客户产品名称", "直接获取"],
        ["生产类型", "直接获取"],
        ["计划数量", "直接获取"],
        ["单支重量", "直接获取"],
        ["定长(mm)", "直接获取"],
        ["挤压机台", "直接获取"],
        ["来源料框", "排程明细获取"],
    ], [1538, 2160])

    add_heading(doc, "3.3 物料清单", 2)
    add_table(doc, ["字段", "方式"], [
        ["序号", "页面计算"],
        ["框号", "来源料框直接获取"],
        ["库位号", "来源料框直接获取"],
        ["是否CPK", "质量记录获取"],
        ["产品名称", "排程直接获取"],
        ["炉次号", "来源料框直接获取"],
        ["挤压批次", "来源料框直接获取"],
        ["挤压机台", "排程直接获取"],
        ["状态", "业务动作更新"],
        ["模具号", "排程直接获取"],
        ["排程编号", "排程直接获取"],
        ["排程类型", "排程直接获取"],
        ["数量", "来源料框直接获取"],
        ["定长(mm)", "排程直接获取"],
        ["净重(kg)", "数量×单支重量"],
        ["上料时间", "确认上料时系统生成"],
        ["完工时间", "确认报工时系统生成"],
        ["良品数量", "报工累计计算"],
        ["不良品数量", "报工累计计算"],
    ], [1538, 2160])

    add_heading(doc, "3.4 收料与裁切报工", 2)
    add_table(doc, ["字段", "方式"], [
        ["报工班组", "手动选择"],
        ["报工时间", "系统生成"],
        ["来料数量", "物料清单直接获取"],
        ["已裁切支数", "良品数量+不良品数量"],
        ["已登记不良长支数", "不良品数量直接获取"],
        ["本次裁切支数", "剩余数量默认/手动调整"],
        ["实收支数", "人工录入（待接通）"],
        ["不良品片数", "人工录入"],
        ["不良品原因", "人工录入（待接通）"],
        ["本次良品数量", "本次裁切支数-不良品片数"],
        ["累计良品数量", "原良品+本次良品"],
        ["累计不良数量", "原不良+本次不良"],
        ["完工状态", "报工规则判断"],
        ["完工时间", "完工时系统生成"],
    ], [1538, 2160])
    add_bullet(doc, "不良品片数不得大于本次裁切支数；本次裁切支数不得超过来料数量减已裁切支数。")
    add_bullet(doc, "当前代码提交报工后直接标记已完工，且“实收支数、不良原因、报工班组”未写入生产记录；正式实现必须补充校验和持久化。")

    add_heading(doc, "3.5 装托与装箱", 2)
    add_table(doc, ["字段", "方式"], [
        ["来源料框", "同排程已收料料框获取"],
        ["栈板编号", "系统生成"],
        ["装托数量", "镭雕码数量汇总"],
        ["栈板最大载量", "人工录入"],
        ["栈板皮重(kg)", "人工录入"],
        ["单支裁切数量", "人工录入"],
        ["是否装箱", "人工选择"],
        ["每箱片数", "人工录入"],
        ["预装箱片数", "人工录入"],
        ["小箱数量", "向上取整(预装箱片数÷每箱片数)"],
        ["小箱码", "栈板编号+两位箱序号"],
        ["镭雕码", "扫码输入"],
        ["长支内部码", "系统生成（原型）"],
        ["镭雕时间", "系统生成"],
        ["长支数量", "扫码明细汇总"],
        ["所属小箱", "扫码时绑定"],
    ], [1538, 2160])
    add_bullet(doc, "同一栈板内镭雕码不得重复；装托数量不得超过栈板最大载量；单箱绑定数量不得超过计划片数。")
    add_bullet(doc, "栈板编号规则为 YYYYMMDD-两位流水号；小箱码规则为 栈板编号-两位箱序号。")

    add_heading(doc, "3.6 标识打印与装托记录", 2)
    add_table(doc, ["字段", "方式"], [
        ["标签类型", "人工选择物料/小箱"],
        ["二维码内容", "栈板号或栈板号-BOX"],
        ["标签排程信息", "排程直接获取"],
        ["标签物料数量", "装托数量直接获取"],
        ["装托记录编号", "系统生成"],
        ["来源料框号", "直接获取"],
        ["栈板号", "直接获取"],
        ["打印状态", "完成装托时更新"],
        ["更新时间", "系统生成"],
        ["班组", "当前上班班组获取"],
        ["长支明细", "扫码记录获取"],
        ["小箱明细", "小箱及绑定记录获取"],
        ["下工序长度(m)", "定长(mm)÷1000"],
        ["是否有码", "长支明细数量>0"],
        ["下工序状态", "固定为料框列表"],
    ], [1538, 2160])

    add_heading(doc, "3.7 当前数据落地边界", 2)
    add_bullet(doc, "已下发裁切排程保存在浏览器 localStorage：mom_issued_cutting_schedules。")
    add_bullet(doc, "完成装托后的下工序记录保存在浏览器 localStorage：mom_cutting_packaging_queue。")
    add_bullet(doc, "上下班状态、物料清单收料状态和裁切报工结果主要保存在当前页面内存，刷新页面后不能保证保留。")
    add_bullet(doc, "物料清单和装托记录仍包含静态演示数据；正式上线必须改为服务端数据库和接口。")
    add_bullet(doc, "“实收支数”和“不良品原因”当前仅在报工弹窗录入，提交后未进入后续记录，不能作为正式统计字段。")

    core = doc.core_properties
    core.title = "裁切工作台需求分析书"
    core.subject = "裁切工作台字段、数据获取方式与计算方式"
    core.author = "Codex"
    core.keywords = "MOM, 裁切工作台, 需求分析, 字段来源, 计算方式"
    core.comments = "依据当前裁切工作台代码及用户提供模板生成"

    doc.save(OUTPUT)
    if sha256(REFERENCE) != EXPECTED_SHA256:
        raise RuntimeError("Reference DOCX was modified unexpectedly")
    print(OUTPUT)


if __name__ == "__main__":
    build()
