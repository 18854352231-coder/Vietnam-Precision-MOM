from __future__ import annotations

import hashlib
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


REFERENCE = Path(r"C:\Users\ptn_z\Desktop\挤压生产报表需求分析书.docx")
OUTPUT = Path(r"D:\越南MOM\MOM\docs\裁切生产报表需求分析书.docx")
EXPECTED_HASH = "e6450c96a87602b130f010963bf4e170f1e5640a54c86eca72f287a716de5a1e"
HEADER_FILL = "F2F4F7"
HEADER_TEXT = "1F4D78"
BODY_TEXT = "202124"
BORDER = "D0D5DD"


def file_hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def clear_body(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def style_run(run, size=11, bold=None, color=BODY_TEXT) -> None:
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_text(cell, value: str, header=False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    if header:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(str(value))
    style_run(run, size=8, bold=header, color=HEADER_TEXT if header else BODY_TEXT)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if header:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), HEADER_FILL)
        cell._tc.get_or_add_tcPr().append(shd)


def set_cell_width(cell, width: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    node = tc_pr.find(qn("w:tcW"))
    if node is None:
        node = OxmlElement("w:tcW")
        tc_pr.append(node)
    node.set(qn("w:w"), str(width))
    node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    total = sum(widths)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    for tag, attrs in (
        ("tblW", {"w": str(total), "type": "dxa"}),
        ("tblInd", {"w": "120", "type": "dxa"}),
        ("tblLayout", {"type": "fixed"}),
    ):
        node = tbl_pr.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tbl_pr.append(node)
        for key, value in attrs.items():
            node.set(qn(f"w:{key}"), value)

    cell_mar = tbl_pr.find(qn("w:tblCellMar"))
    if cell_mar is None:
        cell_mar = OxmlElement("w:tblCellMar")
        tbl_pr.append(cell_mar)
    for side, amount in (("top", 80), ("bottom", 80), ("start", 120), ("end", 120)):
        node = cell_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            cell_mar.append(node)
        node.set(qn("w:w"), str(amount))
        node.set(qn("w:type"), "dxa")

    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), BORDER)

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        node = OxmlElement("w:gridCol")
        node.set(qn("w:w"), str(width))
        grid.append(node)

    for row in table.rows:
        cant_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant_split)
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])


def mark_header(row) -> None:
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(node)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, header=True)
    mark_header(table.rows[0])
    for values in rows:
        row = table.add_row()
        for index, value in enumerate(values):
            set_cell_text(row.cells[index], value)
    set_table_geometry(table, widths)
    doc.add_paragraph()


def heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True


def bullet(doc: Document, text: str) -> None:
    doc.add_paragraph(text, style="List Bullet")


def number(doc: Document, text: str) -> None:
    doc.add_paragraph(text, style="List Number")


def build() -> None:
    if file_hash(REFERENCE) != EXPECTED_HASH:
        raise RuntimeError("Reference DOCX changed; re-distillation required")
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(REFERENCE, OUTPUT)
    doc = Document(OUTPUT)
    clear_body(doc)

    doc.add_paragraph("裁切生产报表需求分析书", style="Title")

    heading(doc, "文档目的与范围", 1)
    doc.add_paragraph(
        "按照裁切工作台已完工料框生成可追溯的裁切生产数据；裁切报表只展示状态为“已完工”的记录，页面、汇总和 Excel 导出必须使用同一数据口径。"
    )

    heading(doc, "1.1 建设目标", 2)
    bullet(doc, "按已完工料框查看排程、来源追溯、裁切数量、质量结果和生产时间，形成裁切生产追溯明细。")
    bullet(doc, "只保留能够从裁切排程、来源料框和报工记录直接取得，或能由可靠明细计算得到的字段。")
    bullet(doc, "明确页面字段、汇总指标和 Excel 导出的统一计算公式，避免同一指标出现不同数值。")

    heading(doc, "1.2 目标数据流", 2)
    number(doc, "裁切排程下发时，以排程编号同步排程主数据及来源料框明细。")
    number(doc, "裁切工作台确认上料时，记录来源料框、上料时间、裁切机台和状态。")
    number(doc, "裁切报工时，按来源料框记录本次裁切数量、良品数量、不良品数量、不良原因、班组和完工时间。")
    number(doc, "报表服务固定筛选状态为已完工的料框，再按查询条件读取排程、来源料框和报工记录。")
    number(doc, "页面、分页、指标汇总和 Excel 导出复用同一查询结果及同一字段口径。")

    heading(doc, "2. 查询条件需求", 1)
    add_table(doc, ["字段", "控件", "查询规则"], [
        ["生产日期", "日期区间", "按完工时间的日期筛选，闭区间查询"],
        ["排程编号", "手动输入", "去首尾空格，不区分大小写，包含匹配"],
        ["来源料框", "手动输入", "去首尾空格，不区分大小写，包含匹配"],
        ["产品名称", "手动输入", "去首尾空格，不区分大小写，包含匹配"],
        ["挤压批次", "手动输入", "去首尾空格，不区分大小写，包含匹配"],
        ["排程类型", "下拉选择", "产发物料、量产物料、重工物料精确匹配"],
    ], [1581, 1113, 5215])
    bullet(doc, "查询区域不提供执行状态条件；服务端固定限定已完工记录。")
    bullet(doc, "点击“查询”后应用条件；点击“重置”清空条件并回到第一页。")
    bullet(doc, "导出全部筛选结果，不受当前分页限制。")

    heading(doc, "3. 明细字段与数据获取方式", 1)

    heading(doc, "3.1 排程与料框", 2)
    add_table(doc, ["字段", "方式"], [
        ["序号", "页面计算：(页码-1)×每页条数+行索引+1"],
        ["生产日期", "由完工时间截取日期"],
        ["排程编号", "裁切排程直接获取"],
        ["排程类型", "裁切排程直接获取"],
        ["来源料框", "来源料框明细直接获取"],
        ["产品名称", "裁切排程直接获取"],
    ], [1538, 2160])

    heading(doc, "3.2 来源追溯", 2)
    add_table(doc, ["字段", "方式"], [
        ["挤压批次", "来源料框明细直接获取"],
        ["炉次号", "来源料框明细直接获取"],
        ["模具号", "裁切排程直接获取"],
        ["锯台", "裁切报工记录的机台直接获取"],
        ["定长(mm)", "裁切排程直接获取"],
        ["来料净重(kg)", "来源称重直接获取；无称重时可按来料支数×单支重量计算"],
    ], [1538, 2160])

    heading(doc, "3.3 裁切实绩", 2)
    add_table(doc, ["字段", "方式"], [
        ["来料支数", "来源料框数量直接获取"],
        ["已裁切", "良品数量+不良品数量"],
        ["剩余未裁", "MAX(来料支数-已裁切, 0)"],
        ["良品", "裁切报工累计良品数量"],
        ["不良品", "裁切报工累计不良品数量"],
        ["良品率", "已裁切>0时：良品÷已裁切×100%，保留1位小数"],
    ], [1538, 2160])
    bullet(doc, "良品数量和不良品数量均应为非负整数，且二者之和不得超过来料支数。")
    bullet(doc, "报表只展示已完工记录，正常情况下剩余未裁应为0；若不为0，应进入数据质量告警。")

    heading(doc, "3.4 时间与状态", 2)
    add_table(doc, ["字段", "方式"], [
        ["上料时间", "裁切工作台确认上料时系统生成"],
        ["完工时间", "裁切工作台确认报工完工时系统生成"],
        ["状态", "报工完成后写入已完工；报表固定只读取已完工"],
    ], [1538, 2160])

    heading(doc, "4. 汇总指标与导出", 1)
    add_table(doc, ["指标", "计算方式"], [
        ["来源料框数", "COUNT(筛选后的已完工料框记录)"],
        ["来料总数", "SUM(来料支数)"],
        ["已裁切总数", "SUM(良品数量)+SUM(不良品数量)"],
        ["良品数", "SUM(良品数量)"],
        ["不良品数", "SUM(不良品数量)"],
        ["综合良品率", "已裁切总数>0时：良品数÷已裁切总数×100%"],
        ["完工料框", "COUNT(状态=已完工)；当前与来源料框数相同"],
        ["来料净重", "SUM(来料净重)，保留1位小数"],
    ], [1538, 2160])
    bullet(doc, "Excel 工作表“指标汇总”导出上述指标；“料框生产明细”导出第3章全部字段。")
    bullet(doc, "页面、汇总和导出均不得包含待收料或已收料记录。")

    heading(doc, "5. 当前数据获取边界", 1)
    bullet(doc, "CuttingReport.vue 当前 sourceRows 为8条静态测试数据，报表未读取裁切工作台生产记录。")
    bullet(doc, "裁切工作台可取得排程、来源料框、上料时间、良品、不良品和完工时间，但物料清单及报工结果主要保存在页面内存。")
    bullet(doc, "锯台目前是工作台当前选择值，没有随每条报工记录稳定保存，正式报表无法可靠追溯。")
    bullet(doc, "生产日期建议统一由完工时间生成；当前报表中的 productionDate 为测试数据独立字段。")
    bullet(doc, "正式上线必须增加裁切生产记录服务端落库和报表查询接口，并以排程编号+来源料框号作为基本关联键。")

    doc.core_properties.title = "裁切生产报表需求分析书"
    doc.core_properties.subject = "裁切生产报表字段与数据获取、计算方式"
    doc.core_properties.author = "Codex"
    doc.core_properties.keywords = "MOM, 裁切报表, 需求分析, 字段来源, 计算公式"
    doc.core_properties.comments = "依据当前裁切报表和裁切工作台代码生成"
    doc.save(OUTPUT)

    if file_hash(REFERENCE) != EXPECTED_HASH:
        raise RuntimeError("Reference DOCX was modified")
    print(OUTPUT)


if __name__ == "__main__":
    build()
